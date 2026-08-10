"""Fixtures compartidos para toda la suite de tests.

Ningún test de esta suite (salvo test_qa_wosg.py, marcado @pytest.mark.live)
llama a la API real de Anthropic: `agents.base.ask` se reemplaza por un doble
programable (ver `scripted_ask` más abajo), así los tests corren rápido,
gratis y de forma determinística.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))


class ScriptedAsk:
    """Doble de test para utils.llm.ask / agents.base.ask.

    Se configura con reglas (predicado sobre el system_prompt -> respuesta) y
    se usa como `monkeypatch.setattr("agents.base.ask", scripted)`. Guarda
    todas las llamadas recibidas en `.calls` para poder inspeccionarlas.
    """

    def __init__(self):
        self.rules: list[tuple[callable, callable]] = []
        self.calls: list[dict] = []

    def when(self, predicate):
        """Decorador: registra una regla. `predicate(system_prompt) -> bool`."""

        def register(response_fn):
            self.rules.append((predicate, response_fn))
            return response_fn

        return register

    def add(self, predicate, response):
        """Registra una regla sin decorador. `response` puede ser un string
        fijo o una función (system_prompt, user_prompt, cached_context) -> str.
        """
        fn = response if callable(response) else (lambda *a, **kw: response)
        self.rules.append((predicate, fn))

    def __call__(self, system_prompt, user_prompt, model=None, max_tokens=None, cached_context=None, tools=None):
        self.calls.append(
            dict(
                system_prompt=system_prompt,
                user_prompt=user_prompt,
                model=model,
                max_tokens=max_tokens,
                cached_context=cached_context,
                tools=tools,
            )
        )
        for predicate, response_fn in self.rules:
            if predicate(system_prompt):
                return response_fn(system_prompt, user_prompt, cached_context)
        raise AssertionError(
            f"ScriptedAsk: ninguna regla matcheó este system_prompt "
            f"(primeros 80 chars): {system_prompt[:80]!r}"
        )


def contains(*substrings: str):
    """Predicado: el system_prompt (normalizado, sin saltos de línea) contiene
    TODAS las substrings dadas. Usar esto en vez de matchear contra el string
    crudo — los prompts multilínea de los agentes envuelven texto en puntos
    donde un match ingenuo se rompe."""

    def predicate(system_prompt: str) -> bool:
        normalized = " ".join(system_prompt.split())
        return all(s in normalized for s in substrings)

    return predicate


@pytest.fixture
def scripted_ask():
    return ScriptedAsk()


def extract_requested_ids(cached_context: str) -> list[int]:
    """Extrae los IDs de pregunta del bloque 'PREGUNTAS A RESPONDER' que el
    Agente 11 arma en el contexto cacheado (formato 'N. [sección] texto')."""
    return [int(m) for m in re.findall(r"^(\d+)\.\s\[", cached_context or "", re.MULTILINE)]


@pytest.fixture
def mini_pabrai_workbook(tmp_path) -> Path:
    """Crea un .xlsx mínimo con la misma estructura que el checklist real
    del usuario (PLANTILLA (duplicar) con headings de sección + preguntas
    numeradas), pero con pocas preguntas para que los tests sean rápidos de
    leer. La estructura (columnas, fila de header, tipos de celda) es
    idéntica a la del archivo real, así el código bajo test es el mismo."""
    path = tmp_path / "checklist_pabrai_test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PLANTILLA (duplicar)"
    ws["A1"] = "CHECKLIST PABRAI — [TICKER — NOMBRE EMPRESA]"
    ws["A6"], ws["B6"], ws["C6"], ws["D6"], ws["E6"] = "#", "PREGUNTA", "VEREDICTO", "SEVERIDAD", "NOTAS"

    rows = [
        ("1. APALANCAMIENTO / DEUDA  ★ SECCIÓN CRÍTICA", None),
        ("Pabrai: cita de contexto, no es un heading de sección.", None),
        (1, "¿Cuál es el ratio Deuda Total / Equity?"),
        (2, "¿Cuál es el ratio Deuda Neta / EBITDA?"),
        (3, "¿Hay vencimientos importantes de deuda en los próximos 2-3 años?"),
        ("2. VENTAJA COMPETITIVA / MOAT  ★ SECCIÓN CRÍTICA", None),
        (4, "¿Tiene ventaja competitiva durable?"),
        (5, "¿El moat se está expandiendo, es estable, o se erosiona?"),
        ("3. MANAGEMENT / PROPIEDAD  ★ SECCIÓN CRÍTICA", None),
        (6, "¿Directivos poseen acciones significativas?"),
        ("4. VALORACIÓN / MARGEN DE SEGURIDAD", None),
        (7, "¿Precio con descuento significativo sobre valor intrínseco?"),
        (8, "¿FCF yield actual atractivo?"),
        ("8. CONTABILIDAD / CALIDAD DE EARNINGS", None),
        (9, "¿Earnings de alta calidad (cash-backed)?"),
    ]
    row_num = 7
    for a_val, b_val in rows:
        ws.cell(row=row_num, column=1, value=a_val)
        if b_val is not None:
            ws.cell(row=row_num, column=2, value=b_val)
        row_num += 1

    wb.save(path)
    return path
