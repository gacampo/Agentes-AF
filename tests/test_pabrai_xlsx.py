"""Tests de utils/pabrai_xlsx.py: lectura del schema de preguntas, escritura
full/lite, y recálculo del panel de evaluación. Usa el fixture
`mini_pabrai_workbook` (estructura idéntica al archivo real, con pocas
preguntas para tests rápidos y sin exponer datos propios del usuario)."""

from __future__ import annotations

import openpyxl

from utils.pabrai_xlsx import (
    apply_full_checklist,
    apply_lite_refresh,
    load_question_schema,
    _panel_rows,
)


def test_load_question_schema_lee_preguntas_secciones_y_criticidad(mini_pabrai_workbook):
    wb = openpyxl.load_workbook(mini_pabrai_workbook)
    questions = load_question_schema(wb)

    assert len(questions) == 9
    assert [q.id for q in questions] == list(range(1, 10))

    by_id = {q.id: q for q in questions}
    assert by_id[1].seccion.startswith("1. APALANCAMIENTO")
    assert by_id[4].seccion.startswith("2. VENTAJA COMPETITIVA")
    assert by_id[6].seccion.startswith("3. MANAGEMENT")
    assert by_id[7].seccion.startswith("4. VALORACIÓN")
    assert by_id[9].seccion.startswith("8. CONTABILIDAD")

    # Las primeras 3 secciones (Leverage, Moat, Management) son las críticas
    assert by_id[1].is_critical and by_id[4].is_critical and by_id[6].is_critical
    assert not by_id[7].is_critical and not by_id[9].is_critical

    # La cita de Pabrai (fila suelta sin patrón "N. ") no debe crear una sección espuria
    secciones = {q.seccion for q in questions}
    assert not any("Pabrai:" in s for s in secciones)


def _answers(ids_veredictos: dict[int, str]) -> dict[int, dict]:
    return {i: {"veredicto": v, "severidad": "Menor", "notas": f"nota {i}"} for i, v in ids_veredictos.items()}


def test_apply_full_checklist_escribe_todas_las_respuestas_y_calcula_panel(mini_pabrai_workbook):
    answers = _answers({i: "OK" for i in range(1, 10)})
    answers[2]["veredicto"] = "⚠️ Red Flag"  # crítica (Leverage)

    resumen = apply_full_checklist(
        str(mini_pabrai_workbook), "TESTCO", "CHECKLIST PABRAI — TESTCO", "Precio: $100", answers, "FULL-2026"
    )

    assert resumen["preguntas_escritas"] == 9
    assert resumen["showstoppers"] == []
    assert resumen["red_flags_criticas"] == 1
    assert resumen["red_flags_secundarias"] == 0
    assert "5%" in resumen["sizing_recomendado"] and "estándar" in resumen["sizing_recomendado"]

    wb2 = openpyxl.load_workbook(mini_pabrai_workbook)
    assert "TESTCO" in wb2.sheetnames
    ws = wb2["TESTCO"]
    assert ws["A1"].value == "CHECKLIST PABRAI — TESTCO"
    assert ws["C10"].value == "⚠️ Red Flag"  # fila de la pregunta 2 (deuda/EBITDA)


def test_apply_full_checklist_es_idempotente_en_la_hoja(mini_pabrai_workbook):
    """Correr FULL dos veces para la misma empresa reusa la hoja, no la duplica."""
    answers = _answers({i: "OK" for i in range(1, 10)})
    apply_full_checklist(str(mini_pabrai_workbook), "TESTCO", "H", "S", answers, "FULL-1")
    apply_full_checklist(str(mini_pabrai_workbook), "TESTCO", "H", "S", answers, "FULL-2")

    wb2 = openpyxl.load_workbook(mini_pabrai_workbook)
    assert wb2.sheetnames.count("TESTCO") == 1


def test_apply_lite_refresh_agrega_bloque_de_columnas_y_preserva_no_tocadas(mini_pabrai_workbook):
    full_answers = _answers({i: "OK" for i in range(1, 10)})
    full_answers[4]["veredicto"] = "⚠️ Red Flag"  # Moat, crítica — no se toca en el refresh
    apply_full_checklist(str(mini_pabrai_workbook), "TESTCO", "H", "S", full_answers, "FULL-2026")

    # Refresh: solo Apalancamiento (1-3) — ahora con un showstopper nuevo en la pregunta 2
    lite_answers = _answers({1: "OK", 2: "🛑 Showstopper", 3: "OK"})
    resumen = apply_lite_refresh(str(mini_pabrai_workbook), "TESTCO", lite_answers, "Q3-2026")

    assert resumen["preguntas_escritas"] == 3
    assert resumen["showstoppers"] == [2]
    # La red flag de Moat (pregunta 4, no tocada en el refresh) debe seguir contando
    assert resumen["red_flags_criticas"] == 1
    assert "NO INVERTIR" in resumen["sizing_recomendado"]

    wb2 = openpyxl.load_workbook(mini_pabrai_workbook)
    ws = wb2["TESTCO"]
    header = [ws.cell(row=6, column=c).value for c in range(3, 9)]
    assert header[0] == "VEREDICTO FULL-2026"
    assert header[3] == "VEREDICTO Q3-2026"
    # La pregunta 4 (moat) no se tocó en el refresh -> su bloque nuevo queda vacío
    assert ws.cell(row=13, column=6).value in (None, "")  # fila de la pregunta 4, columna F (bloque refresh)


def test_apply_lite_refresh_sin_full_previo_falla_con_error_claro(mini_pabrai_workbook):
    import pytest

    with pytest.raises(ValueError, match="corré primero"):
        apply_lite_refresh(str(mini_pabrai_workbook), "NOEXISTE", _answers({1: "OK"}), "Q1-2026")


def test_historial_de_sizing_acumula_una_fila_por_corrida(mini_pabrai_workbook):
    # Las filas del panel/historial ya NO son un número fijo (ver
    # utils/pabrai_xlsx._panel_rows): se calculan a partir de la última fila
    # de pregunta real del template, para no romper si se agrega/quita una
    # sección de preguntas (como pasó con "12. GOBIERNO CORPORATIVO"). Este
    # test recalcula esa misma fila en vez de hardcodearla, así sigue siendo
    # una prueba real del comportamiento y no un valor mágico congelado.
    wb0 = openpyxl.load_workbook(mini_pabrai_workbook)
    questions0 = load_question_schema(wb0)
    history_header_row = _panel_rows(questions0)["history_header"]

    answers = _answers({i: "OK" for i in range(1, 10)})
    apply_full_checklist(str(mini_pabrai_workbook), "TESTCO", "H", "S", answers, "FULL-2026")
    apply_lite_refresh(str(mini_pabrai_workbook), "TESTCO", _answers({1: "OK"}), "Q1-2026")
    apply_lite_refresh(str(mini_pabrai_workbook), "TESTCO", _answers({1: "OK"}), "Q2-2026")

    wb2 = openpyxl.load_workbook(mini_pabrai_workbook)
    ws = wb2["TESTCO"]
    data_start = history_header_row + 1
    fechas = [ws.cell(row=r, column=1).value for r in range(data_start, data_start + 4)]
    fechas = [f for f in fechas if f]
    assert fechas == ["FULL-2026", "Q1-2026", "Q2-2026"]
